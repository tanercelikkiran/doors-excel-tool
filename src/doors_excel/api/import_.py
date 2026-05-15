"""Import API — orchestrates Excel→DOORS import pipeline."""
from __future__ import annotations

import sqlite3
from pathlib import Path
from typing import TYPE_CHECKING

from loguru import logger

from doors_excel.api.sessions import SessionManager
from doors_excel.api.staging import load_excel_to_staging
from doors_excel.api.diff import run_diff
from doors_excel.common.types import ConflictPolicy
from doors_excel.core.diff.conflict import apply_conflict_policy as _apply_policy
from doors_excel.core.diff.summary import DiffSummary
from doors_excel.core.transformation.hashing import hash_markdown as _hash_md
from doors_excel.core.transformation.markdown_to_rtf import markdown_to_rtf
from doors_excel.core.transformation.rtf_to_markdown import rtf_to_markdown
from doors_excel.core.validation.models import ModuleConfig, ProjectConfig
from doors_excel.infrastructure.database.repositories import (
    StagingBaselineRepository,
    StagingDoorsRepository,
)
from doors_excel.infrastructure.doors.exporter import DoorsExporter
from doors_excel.infrastructure.doors.importer import DoorsImporter
from doors_excel.infrastructure.excel.reader import FormulaPolicy, open_workbook

if TYPE_CHECKING:
    from openpyxl import Workbook
    from openpyxl.worksheet.worksheet import Worksheet


def resolve_worksheet_module(
    wb: Workbook,
    ws: Worksheet,
    project_cfg: ProjectConfig,
) -> ModuleConfig | None:
    """Determine which ModuleConfig corresponds to *ws*.

    Priority:
    1. Custom workbook property _DOORS_Module_Path::{sheet.title} matching module_path
    2. Sheet name matches make_sheet_name(module_name, module_path)
    3. Case-insensitive module name match against sheet title
    """
    from doors_excel.infrastructure.excel.metadata import get_module_path
    from doors_excel.infrastructure.excel.naming import make_sheet_name

    stored_path = get_module_path(wb, ws)
    if stored_path:
        match = next((m for m in project_cfg.modules if m.module_path == stored_path), None)
        if match:
            return match

    for mod_cfg in project_cfg.modules:
        mod_name = mod_cfg.module_path.rstrip("/").rsplit("/", 1)[-1]
        if make_sheet_name(mod_name, mod_cfg.module_path) == ws.title:
            return mod_cfg

    ws_title_lower = ws.title.lower()
    for mod_cfg in project_cfg.modules:
        mod_name = mod_cfg.module_path.rstrip("/").rsplit("/", 1)[-1].lower()
        if mod_name == ws_title_lower:
            return mod_cfg

    return None


def bulk_stage_imports(
    excel_path: Path | str,
    project_cfg: ProjectConfig,
    *,
    db_path: Path | str,
    doors_conn: object,
    trim_whitespace: bool = True,
) -> list[tuple[str, str, DiffSummary, ModuleConfig]]:
    """Stage all matched worksheets in *excel_path*.

    Returns list of (session_id, sheet_title, diff_summary, module_config).
    Sheets that don't match any ModuleConfig are silently skipped.
    """
    import openpyxl as _openpyxl

    wb = _openpyxl.load_workbook(excel_path, data_only=True)
    results = []
    for ws in wb.worksheets:
        mod_cfg = resolve_worksheet_module(wb, ws, project_cfg)
        if mod_cfg is None:
            continue
        session_id, stats = stage_import(
            excel_path,
            mod_cfg,
            db_path=db_path,
            doors_conn=doors_conn,
            trim_whitespace=trim_whitespace,
            sheet_title=ws.title,
        )
        results.append((session_id, ws.title, stats, mod_cfg))
    return results


def stage_import(
    excel_path: Path | str,
    module_config: ModuleConfig,
    *,
    db_path: Path | str,
    doors_conn: object,
    baseline: str = "current",
    trim_whitespace: bool = True,
    sheet_title: str | None = None,
) -> tuple[str, DiffSummary]:
    """Stage an Excel import: load Excel + DOORS data into SQLite, compute diff.

    Returns ``(session_id, DiffSummary)``.
    """
    from doors_excel.api.validate import _pick_worksheet

    p = Path(excel_path)
    attributes = [m.attribute for m in module_config.column_mappings]

    mgr = SessionManager(db_path)
    try:
        info = mgr.create(p, module_config.module_path)
        sid = info.session_id
        conn = mgr.conn

        # --- Stage Excel ---
        wb = open_workbook(p, formula_policy=FormulaPolicy.DATA_ONLY)
        ws = wb[sheet_title] if sheet_title is not None else _pick_worksheet(wb, module_config)
        load_excel_to_staging(ws, conn, sid, module_config, trim_whitespace=trim_whitespace)

        # --- Stage DOORS (current state = staging_doors AND staging_baseline) ---
        exporter = DoorsExporter(doors_conn)
        raw_rows = exporter.export_module(module_config.module_path, attributes, baseline=baseline)

        text_attrs = {
            m.attribute for m in module_config.column_mappings if m.attribute_type == "Text"
        }
        for row in raw_rows:
            if row["attribute"] in text_attrs and row.get("rtf_value"):
                result = rtf_to_markdown(row["rtf_value"])
                row["md_hash"] = _hash_md(result.markdown)

        StagingDoorsRepository(conn).insert_many(
            [{**row, "session_id": sid} for row in raw_rows]
        )
        StagingBaselineRepository(conn).insert_many(
            [
                {
                    "session_id": sid,
                    "object_id": row["object_id"],
                    "attribute": row["attribute"],
                    "value": row["value"],
                    "object_type": row["object_type"],
                    "level": row["level"],
                    "parent_id": row["parent_id"],
                }
                for row in raw_rows
            ]
        )

        summary = run_diff(conn, sid)
    finally:
        mgr.close()

    return sid, summary


def execute_import(
    session_id: str,
    conn: sqlite3.Connection,
    *,
    doors_conn: object,
    conflict_policy: ConflictPolicy = "excel-wins",
    module_path: str | None = None,
    module_config: "ModuleConfig | None" = None,
    include_new: bool = False,
    deletion_policy: str = "ignore",
    soft_delete_attribute: str = "Status",
    soft_delete_value: str = "Deleted",
    accept_ole_overwrites: bool = False,
) -> int:
    """Apply UPDATED and CONFLICT diff_results to DOORS. Returns count applied.

    Text-type attributes are converted from GFM Markdown to RTF before writing.
    If the md_hash is unchanged (REQ-FUN-105.4) and DOORS value matches baseline,
    the original RTF is restored from rollback_snapshots to avoid formatting loss.

    When *include_new* is True, rows with change_type='NEW' are also created as
    new DOORS objects via :meth:`DoorsImporter.create_objects`.
    """
    _apply_policy(conn, session_id, conflict_policy)

    rows = conn.execute(
        """
        SELECT dr.object_id, dr.attribute,
               CASE WHEN dr.change_type = 'CONFLICT' THEN dr.resolved_value
                    ELSE dr.excel_value
               END AS apply_value,
               s.doors_module,
               se.md_hash  AS excel_md_hash,
               sd.md_hash  AS doors_md_hash,
               dr.doors_value,
               dr.baseline_value,
               rs.original_rtf
        FROM diff_results dr
        JOIN sessions s ON s.session_id = dr.session_id
        LEFT JOIN (
            SELECT session_id, object_id, attribute, md_hash
            FROM staging_excel
            WHERE session_id = :sid
            GROUP BY session_id, object_id, attribute
        ) se ON se.object_id = dr.object_id AND se.attribute = dr.attribute
        LEFT JOIN staging_doors sd
            ON sd.session_id = dr.session_id
           AND sd.object_id  = dr.object_id
           AND sd.attribute  = dr.attribute
        LEFT JOIN rollback_snapshots rs
            ON rs.session_id = dr.session_id
           AND rs.object_id  = dr.object_id
           AND rs.attribute  = dr.attribute
        WHERE dr.session_id = :sid
          AND (
              (dr.change_type = 'UPDATED' AND dr.attribute != '_Parent_ID')
              OR (dr.change_type = 'CONFLICT' AND dr.resolved_value IS NOT NULL
                  AND dr.attribute != '_Parent_ID')
          )
        """,
        {"sid": session_id},
    ).fetchall()

    # Filter out OLE-protected objects if accept_ole_overwrites is False
    if not accept_ole_overwrites and rows:
        row_object_ids = list({r["object_id"] for r in rows})
        if row_object_ids:
            placeholders = ",".join("?" * len(row_object_ids))
            ole_rows = conn.execute(
                f"SELECT DISTINCT object_id FROM staging_doors "
                f"WHERE session_id = ? AND has_ole = 1 AND object_id IN ({placeholders})",
                [session_id] + row_object_ids,
            ).fetchall()
            ole_ids = {r["object_id"] for r in ole_rows}
            if ole_ids:
                rows = [r for r in rows if r["object_id"] not in ole_ids]
                logger.warning(
                    "Skipped {} object(s) with embedded OLE content. "
                    "Use --accept-ole-overwrites to allow updates.",
                    len(ole_ids),
                )

    # Resolve module path (param → rows[0] → sessions lookup)
    if module_path is not None:
        mod_path = module_path
    elif rows:
        mod_path = rows[0]["doors_module"]
    else:
        sr = conn.execute(
            "SELECT doors_module FROM sessions WHERE session_id = ?", (session_id,)
        ).fetchone()
        mod_path = sr["doors_module"] if sr else ""

    applied = 0

    if rows:
        text_attrs: set[str] = set()
        if module_config is not None:
            text_attrs = {m.attribute for m in module_config.column_mappings if m.attribute_type == "Text"}

        updates = []
        for r in rows:
            value = r["apply_value"] or ""
            if r["attribute"] in text_attrs:
                excel_hash = r["excel_md_hash"]
                doors_hash = r["doors_md_hash"]
                doors_val = r["doors_value"]
                base_val = r["baseline_value"]
                if (excel_hash is not None
                        and doors_hash is not None
                        and excel_hash == doors_hash
                        and doors_val == base_val
                        and r["original_rtf"] is not None):
                    value = r["original_rtf"]
                else:
                    value = markdown_to_rtf(value)
            updates.append({
                "object_id": r["object_id"],
                "attribute": r["attribute"],
                "value": value,
            })

        importer = DoorsImporter(doors_conn)
        importer.apply_updates(mod_path, updates)
        applied += len(updates)

    # --- Handle MOVED objects ---
    moved_rows = conn.execute(
        """
        SELECT DISTINCT dr.object_id,
               se_parent.value AS new_parent_id,
               se_place.value  AS placement
        FROM diff_results dr
        LEFT JOIN staging_excel se_parent
               ON se_parent.session_id = dr.session_id
              AND se_parent.object_id  = dr.object_id
              AND se_parent.attribute  = '_Parent_ID'
        LEFT JOIN staging_excel se_place
               ON se_place.session_id = dr.session_id
              AND se_place.object_id  = dr.object_id
              AND se_place.attribute  = '_Placement'
        WHERE dr.session_id = :sid
          AND dr.change_type = 'MOVED'
        """,
        {"sid": session_id},
    ).fetchall()

    if moved_rows:
        moves = []
        for r in moved_rows:
            try:
                parent_id = int(r["new_parent_id"]) if r["new_parent_id"] else None
            except (ValueError, TypeError):
                parent_id = None
            moves.append({
                "object_id": r["object_id"],
                "new_parent_id": parent_id,
                "placement": (r["placement"] or "below").lower(),
            })
        DoorsImporter(doors_conn).move_objects(mod_path, moves)
        applied += len(moves)

    if include_new:
        new_objects = _collect_new_objects(conn, session_id, module_config)
        if new_objects:
            DoorsImporter(doors_conn).create_objects(mod_path, new_objects)
            applied += len(new_objects)

    applied += _handle_deleted_objects(
        conn, session_id, mod_path, doors_conn,
        deletion_policy, soft_delete_attribute, soft_delete_value,
    )

    return applied


def _handle_deleted_objects(
    conn: sqlite3.Connection,
    session_id: str,
    mod_path: str,
    doors_conn: object,
    deletion_policy: str,
    soft_delete_attribute: str,
    soft_delete_value: str,
) -> int:
    if deletion_policy == "ignore":
        return 0

    deleted_rows = conn.execute(
        "SELECT DISTINCT object_id FROM diff_results WHERE session_id = ? AND change_type = 'DELETED'",
        (session_id,),
    ).fetchall()
    if not deleted_rows:
        return 0

    object_ids = [r["object_id"] for r in deleted_rows if r["object_id"] is not None]
    importer = DoorsImporter(doors_conn)

    if deletion_policy == "purge":
        importer.delete_objects(mod_path, object_ids)
        return len(object_ids)

    if deletion_policy == "soft-delete":
        updates = [
            {"object_id": oid, "attribute": soft_delete_attribute, "value": soft_delete_value}
            for oid in object_ids
        ]
        importer.apply_updates(mod_path, updates)
        return len(object_ids)

    raise ValueError(
        f"Unknown deletion_policy {deletion_policy!r}. "
        "Expected 'ignore', 'soft-delete', or 'purge'."
    )


_NEW_SKIP_COLS = frozenset({"Absolute Number", "_Parent_ID", "_Placement", "Level", "Parent Absolute Number"})


def _collect_new_objects(
    conn: sqlite3.Connection,
    session_id: str,
    module_config: "ModuleConfig | None" = None,
) -> list[dict]:
    """Return a list of new-object dicts grouped by row_number."""
    new_rows = conn.execute(
        "SELECT DISTINCT row_number FROM diff_results WHERE session_id = ? AND change_type = 'NEW'",
        (session_id,),
    ).fetchall()

    skip_cols = set(_NEW_SKIP_COLS)
    if module_config is not None:
        skip_cols.add(module_config.object_id_column)

    objects = []
    for nr in new_rows:
        rn = nr["row_number"]
        attr_rows = conn.execute(
            "SELECT attribute, value FROM staging_excel WHERE session_id = ? AND row_number = ?",
            (session_id, rn),
        ).fetchall()

        parent_id: int | None = None
        placement: str = "below"  # default: create as child
        attrs: dict[str, str] = {}
        for ar in attr_rows:
            attr, val = ar["attribute"], ar["value"]
            if attr == "_Parent_ID":
                try:
                    parent_id = int(val) if val else None
                except (ValueError, TypeError):
                    parent_id = None
            elif attr == "_Placement":
                placement = (val or "below").lower()
            elif attr not in skip_cols and val is not None:
                attrs[attr] = val

        if attrs:
            objects.append({"parent_id": parent_id, "placement": placement, "attributes": attrs})

    return objects
